from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR

path = r"D:\study\system\low-carbon-dormitory\docs\enterprise-year-end-stage3plus.pptx"
prs = Presentation(path)

GREEN = RGBColor(0x35, 0xB0, 0x7A)
DARK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x63, 0x72, 0x7D)
LIGHT_BG = RGBColor(0xF4, 0xFA, 0xF7)
BORDER = RGBColor(0xC9, 0xE7, 0xD8)
ACCENT_BG = RGBColor(0xE7, 0xF6, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_box(slide, x, y, w, h, fill, line=None, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(radius, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=DARK, font="微软雅黑", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_paragraphs(box, items, size=13, color=DARK, font="微软雅黑", space_after=4):
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = item
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        p.space_after = Pt(space_after)
    return box


def remove_shapes_from(slide, start_index_zero_based):
    for shape in list(slide.shapes)[start_index_zero_based:]:
        sp = shape._element
        sp.getparent().remove(sp)


# 第26页：未来规划
slide = prs.slides[25]
slide.shapes[0].text = "未来规划"
remove_shapes_from(slide, 3)

add_text(slide, Inches(0.78), Inches(0.95), Inches(2.9), Inches(0.4), "AI 时间线与低碳趋势双轮驱动", size=14, color=MUTED)

add_box(slide, Inches(0.75), Inches(1.45), Inches(4.35), Inches(5.2), LIGHT_BG, BORDER)
add_text(slide, Inches(1.0), Inches(1.7), Inches(3.3), Inches(0.35), "AI 从兴起走向校园落地", size=20, bold=True)
add_text(slide, Inches(1.0), Inches(2.05), Inches(3.8), Inches(0.45), "从规则计算到生成式智能，AI 正进入校园治理与行为引导阶段。", size=11, color=MUTED)

line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.2), Inches(2.65), Inches(0.05), Inches(3.2))
line.fill.solid()
line.fill.fore_color.rgb = GREEN
line.line.color.rgb = GREEN

timeline = [
    ("1956", "AI 概念提出", "人工智能成为独立研究方向，完成从“能不能做”到“可以研究”的起点。"),
    ("2012", "深度学习突破", "图像识别等任务精度跃升，AI 从实验室走向大规模场景应用。"),
    ("2022", "生成式 AI 爆发", "大模型推动自然语言交互、知识问答和智能助手快速普及。"),
    ("未来", "校园 AI 智能体", "结合水电、餐损、出行数据做预测预警、节能推荐与个性化积分激励。"),
]

for i, (year, title, desc) in enumerate(timeline):
    y = 2.55 + i * 0.82
    dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.05), Inches(y), Inches(0.35), Inches(0.35))
    dot.fill.solid()
    dot.fill.fore_color.rgb = WHITE
    dot.line.color.rgb = GREEN
    dot.line.width = Pt(2)
    add_text(slide, Inches(1.55), Inches(y - 0.02), Inches(0.7), Inches(0.25), year, size=16, bold=True, color=GREEN)
    add_text(slide, Inches(2.2), Inches(y - 0.03), Inches(2.0), Inches(0.28), title, size=15, bold=True)
    add_text(slide, Inches(2.2), Inches(y + 0.23), Inches(2.55), Inches(0.42), desc, size=10.5, color=MUTED)

add_text(slide, Inches(5.35), Inches(1.68), Inches(3.3), Inches(0.35), "项目下一阶段重点", size=20, bold=True)
plan_cards = [
    ("01 接入出行端", "把手机端低碳出行并入统一平台，打通账号、积分、排行榜与碳画像。"),
    ("02 AI 预测预警", "基于历史水电、餐损、天气与作息数据，识别异常浪费并提前提醒。"),
    ("03 个性化推荐", "按学生、宿舍、楼栋生成节能建议，推动步行骑行、错峰用电与光盘行动。"),
    ("04 运营闭环", "联动后勤、餐厅、宿管与活动运营，让监测、分析、激励、复盘形成长期机制。"),
]
card_y = [2.05, 3.15, 4.25, 5.35]
for (title, desc), y in zip(plan_cards, card_y):
    add_box(slide, Inches(5.3), Inches(y), Inches(4.15), Inches(0.88), ACCENT_BG, BORDER)
    add_text(slide, Inches(5.55), Inches(y + 0.12), Inches(1.8), Inches(0.22), title, size=14, bold=True, color=GREEN)
    add_text(slide, Inches(6.95), Inches(y + 0.1), Inches(2.2), Inches(0.45), desc, size=10.8, color=DARK)

add_text(slide, Inches(5.35), Inches(6.35), Inches(4.2), Inches(0.36), "低碳趋势判断：政策长期化、治理精细化、行为可量化、校园品牌绿色化。", size=11.5, color=MUTED)


# 第27页：总结
slide = prs.slides[26]
slide.shapes[0].text = "总结"
remove_shapes_from(slide, 3)

add_text(slide, Inches(0.78), Inches(0.95), Inches(3.1), Inches(0.4), "从单点功能走向统一低碳校园平台", size=14, color=MUTED)

summary_cards = [
    ("场景整合", "已覆盖宿舍水电、餐厅餐损，并预留低碳出行接入路径。"),
    ("数据闭环", "把记录、统计、分析、积分激励串成可持续运营的管理链路。"),
    ("发展方向", "AI 提升预测与决策能力，低碳趋势提供长期建设目标。"),
]

for i, (title, desc) in enumerate(summary_cards):
    x = 0.85 + i * 3.05
    add_box(slide, Inches(x), Inches(1.65), Inches(2.7), Inches(1.35), WHITE, BORDER)
    add_text(slide, Inches(x + 0.18), Inches(1.88), Inches(1.2), Inches(0.25), title, size=18, bold=True, color=GREEN)
    add_text(slide, Inches(x + 0.18), Inches(2.22), Inches(2.3), Inches(0.58), desc, size=11.5, color=DARK)

add_box(slide, Inches(0.85), Inches(3.35), Inches(8.55), Inches(2.25), LIGHT_BG, BORDER)
add_text(slide, Inches(1.08), Inches(3.62), Inches(1.7), Inches(0.28), "核心结论", size=21, bold=True)
conclusion_box = slide.shapes.add_textbox(Inches(1.1), Inches(4.02), Inches(8.0), Inches(1.25))
add_paragraphs(
    conclusion_box,
    [
        "本项目以校园真实场景为切口，让水电、餐损、出行等低碳行为实现可记录、可计算、可激励。",
        "面向未来，AI 将把系统能力从“事后统计”推进到“实时预测 + 主动引导 + 精准运营”。",
        "随着双碳目标持续推进，高校低碳平台将不只是管理工具，更会成为绿色校园建设与学生价值引导的重要基础设施。",
    ],
    size=13,
    color=DARK,
)

add_box(slide, Inches(0.85), Inches(5.95), Inches(8.55), Inches(0.65), GREEN, GREEN)
add_text(
    slide,
    Inches(1.08),
    Inches(6.12),
    Inches(8.0),
    Inches(0.25),
    "绿色低碳智行系统的下一步，不只是扩功能，而是形成“AI 赋能 + 低碳治理 + 学生参与”的校园新生态。",
    size=14,
    bold=True,
    color=WHITE,
)

prs.save(path)
print(path)
